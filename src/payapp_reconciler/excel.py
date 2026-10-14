"""Excel review workbook: Summary, Line Review (problem cells highlighted and commented), Findings.

The workbook only presents numbers computed elsewhere. It holds values, not formulas, so what the
reviewer sees is exactly what the tool computed.
"""

from __future__ import annotations

from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .loaders import ReviewInput
from .models import Finding, G703Row, ReviewResult, Severity

CURRENCY = '"$"#,##0.00_);[Red]("$"#,##0.00)'
PERCENT = '0.00"%"'  # values are in percent units (45.23 means 45.23%)
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_BORDER = Border(top=Side(style="thin"), bottom=Side(style="double"))
SEVERITY_FILL = {
    Severity.ERROR: PatternFill("solid", fgColor="FFC7CE"),
    Severity.WARNING: PatternFill("solid", fgColor="FFEB9C"),
    Severity.INFO: PatternFill("solid", fgColor="DDEBF7"),
}
AUTHOR = "payapp-reconciler"

# Line Review layout: G703 columns A-I as submitted, then the certifiable columns.
LINE_COLUMNS: list[tuple[str, str, str, int]] = [
    # (finding column key, header, field on G703Row, width)
    ("A", "Item", "item_no", 9),
    ("B", "Description of Work", "description", 38),
    ("C", "Scheduled Value (C)", "scheduled_value", 16),
    ("D", "From Previous (D)", "from_previous", 16),
    ("E", "This Period (E)", "this_period", 16),
    ("F", "Materials Stored (F)", "materials_stored", 16),
    ("G", "Completed & Stored (G)", "total_completed_stored", 18),
    ("pct", "% (G / C)", "pct_complete", 10),
    ("H", "Balance to Finish (H)", "balance_to_finish", 16),
    ("I", "Retainage (I)", "retainage", 14),
]
CERT_HEADERS = [("Certifiable G", 16), ("Certifiable Retainage", 16), ("G Over Certifiable", 16)]
MONEY_KEYS = {"C", "D", "E", "F", "G", "H", "I"}


def _header(ws: Worksheet, row: int, headers: list[tuple[str, int]]) -> None:
    for col, (text, width) in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[row].height = 32


def _flag(cell: Cell, findings: list[Finding]) -> None:
    """Fill a cell by its most severe finding and attach every finding as a comment."""
    worst = min(findings, key=lambda f: list(Severity).index(f.severity))
    cell.fill = SEVERITY_FILL[worst.severity]
    text = "\n\n".join(
        f"{f.rule_id} {f.rule_name} ({f.severity.value}): {f.message}" for f in findings
    )
    cell.comment = Comment(text, AUTHOR, width=420, height=160)


def _summary(ws: Worksheet, result: ReviewResult) -> None:
    ws.title = "Summary"
    ws["A1"] = "Pay Application Review"
    ws["A1"].font = Font(bold=True, size=14)
    info = [
        ("Project", result.project),
        ("General Contractor", result.general_contractor),
        ("Subcontractor", f"{result.sub_name} ({result.sub_id})"),
        ("Trade", result.trade),
        ("Application No.", result.application_no),
        ("Period To", result.period_to),
        ("Certifiable % Complete", result.certifiable_pct_complete),
        ("Retainage Rate", result.retainage_rate),
    ]
    for offset, (label, value) in enumerate(info, start=3):
        ws.cell(row=offset, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=offset, column=2, value=value)
        cell.alignment = Alignment(horizontal="left")
        if label.endswith(("Complete", "Rate")):
            cell.number_format = PERCENT
        elif label == "Period To":
            cell.number_format = "yyyy-mm-dd"

    top = 12
    _header(ws, top, [("Line", 24), ("G702", 44), ("Claimed", 18), ("Certifiable", 18),
                      ("Variance", 18)])  # fmt: skip
    g702_findings: dict[int, list[Finding]] = defaultdict(list)
    for f in result.findings:
        if f.column and f.column.startswith("L"):
            g702_findings[int(f.column[1:])].append(f)
    for i, line in enumerate(result.g702, start=top + 1):
        ws.cell(row=i, column=1, value=line.line)
        ws.cell(row=i, column=2, value=line.label)
        for col, value in ((3, line.claimed), (4, line.certifiable), (5, line.variance)):
            ws.cell(row=i, column=col, value=value).number_format = CURRENCY
        if line.line in g702_findings:
            _flag(ws.cell(row=i, column=3), g702_findings[line.line])
        if line.line == 8:
            for col in range(1, 6):
                ws.cell(row=i, column=col).font = Font(bold=True)

    row = top + len(result.g702) + 2
    ws.cell(row=row, column=1, value="Findings").font = Font(bold=True)
    for offset, severity in enumerate(Severity, start=1):
        ws.cell(row=row + offset, column=1, value=severity.value.title())
        ws.cell(row=row + offset, column=2, value=result.count(severity)).alignment = Alignment(
            horizontal="left"
        )
    note = row + len(Severity) + 2
    explanation = (
        "Variance = claimed - certifiable. Positive means the subcontractor asked for more "
        "than the GC can certify."
    )
    ws.cell(row=note, column=1, value=explanation).font = Font(italic=True)
    ws.freeze_panes = "A2"


def _line_review(ws: Worksheet, review: ReviewInput, result: ReviewResult) -> None:
    headers = [(h, w) for _, h, _, w in LINE_COLUMNS] + CERT_HEADERS
    _header(ws, 1, headers)
    col_of = {key: i for i, (key, *_) in enumerate(LINE_COLUMNS, start=1)}
    cert = {line.item_no: line for line in result.certifiable_lines}

    rows_of: dict[str, list[int]] = defaultdict(list)
    r = 1
    for r, g703 in enumerate(review.submitted, start=2):
        rows_of[g703.item_no].append(r)
        for key, _, field, _ in LINE_COLUMNS:
            cell = ws.cell(row=r, column=col_of[key], value=getattr(g703, field))
            if key in MONEY_KEYS:
                cell.number_format = CURRENCY
            elif key == "pct":
                cell.number_format = PERCENT
        line = cert.get(g703.item_no)
        if line is not None and rows_of[g703.item_no] == [r]:  # duplicates show certifiable once
            _cert_cells(ws, r, line.total_completed_stored, line.retainage, g703)

    # Lines the sub left off the G703 still carry previously certified work.
    submitted_items = {g.item_no for g in review.submitted}
    for line in result.certifiable_lines:
        if line.item_no not in submitted_items:
            r += 1
            rows_of[line.item_no].append(r)
            ws.cell(row=r, column=1, value=line.item_no)
            ws.cell(row=r, column=2, value=f"{line.description} (not on submitted G703)")
            _cert_cells(ws, r, line.total_completed_stored, line.retainage, None)

    total_row = r + 1
    ws.cell(row=total_row, column=2, value="Totals").font = Font(bold=True)
    sums = [col_of[k] for k in ("C", "D", "E", "F", "G", "H", "I")]
    sums += [len(LINE_COLUMNS) + i for i in (1, 2, 3)]
    for col in sums:
        total = sum(
            (ws.cell(row=x, column=col).value or Decimal(0) for x in range(2, total_row)),
            Decimal(0),
        )
        cell = ws.cell(row=total_row, column=col, value=total)
        cell.number_format, cell.font, cell.border = CURRENCY, Font(bold=True), TOTAL_BORDER

    # Highlight and comment the cells each finding points at.
    flagged: dict[tuple[int, int], list[Finding]] = defaultdict(list)
    for f in result.findings:
        if f.column not in col_of:
            continue  # G702 findings live on the Summary sheet
        col = col_of[f.column]
        rows = rows_of.get(f.item_no, []) if f.item_no else list(range(2, total_row))
        for x in rows:
            flagged[(x, col)].append(f)
    for (x, col), found in flagged.items():
        _flag(ws.cell(row=x, column=col), found)

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{total_row - 1}"


def _cert_cells(
    ws: Worksheet, r: int, g: Decimal, retainage: Decimal, submitted: G703Row | None
) -> None:
    claimed_g = submitted.total_completed_stored if submitted else Decimal(0)
    base = len(LINE_COLUMNS)
    for offset, value in ((1, g), (2, retainage), (3, claimed_g - g)):
        ws.cell(row=r, column=base + offset, value=value).number_format = CURRENCY


def _findings(ws: Worksheet, result: ReviewResult) -> None:
    _header(ws, 1, [("#", 5), ("Rule", 7), ("Check", 24), ("Severity", 10), ("Item", 9),
                    ("Column / Line", 13), ("Finding", 90), ("Expected", 16), ("Actual", 16),
                    ("Dollar Impact", 16)])  # fmt: skip
    for r, f in enumerate(result.findings, start=2):
        values = [r - 1, f.rule_id, f.rule_name, f.severity.value, f.item_no or "",
                  f.column or "", f.message, f.expected, f.actual, f.dollar_impact]  # fmt: skip
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=col == 7)
        ws.cell(row=r, column=4).fill = SEVERITY_FILL[f.severity]
        is_pct = f.column == "pct"
        for col in (8, 9):
            ws.cell(row=r, column=col).number_format = PERCENT if is_pct else CURRENCY
        ws.cell(row=r, column=10).number_format = CURRENCY
    if not result.findings:
        ws.cell(row=2, column=7, value="No findings. The application ties out.")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{max(len(result.findings) + 1, 2)}"


def write_workbook(review: ReviewInput, result: ReviewResult, path: Path) -> None:
    wb = Workbook()
    _summary(wb.active, result)
    _line_review(wb.create_sheet("Line Review"), review, result)
    _findings(wb.create_sheet("Findings"), result)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
