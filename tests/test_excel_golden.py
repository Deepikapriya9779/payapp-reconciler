"""Golden-file test for the review workbook's key cells.

Refresh after an intentional change:  UPDATE_GOLDEN=1 pytest tests/test_excel_golden.py
"""

from __future__ import annotations

import json
import os
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from payapp_reconciler.excel import write_workbook
from payapp_reconciler.review import review_application

ROOT = Path(__file__).parents[1]
SAMPLE = ROOT / "data" / "sample"
GOLDEN = ROOT / "tests" / "golden"


def _cell(ws: Worksheet, coord: str) -> dict[str, Any]:
    c = ws[coord]
    out: dict[str, Any] = {"value": None if c.value is None else str(c.value)}
    if c.number_format != "General":
        out["format"] = c.number_format
    if c.fill.fill_type:
        out["fill"] = c.fill.fgColor.rgb
    if c.comment:
        out["comment"] = c.comment.text.split(":", 1)[0]  # rule id, name and severity
    return out


def snapshot(path: Path) -> dict[str, Any]:
    wb = load_workbook(path)
    summary, lines, findings = wb["Summary"], wb["Line Review"], wb["Findings"]
    keys = {
        "Summary": ["A1", *[f"B{r}" for r in range(3, 11)]]
        + [f"{c}{r}" for r in range(13, 22) for c in "ABCDE"],
        "Line Review": [f"{c}1" for c in "ABCDEFGHIJKLM"]
        + [c.coordinate for row in lines.iter_rows() for c in row if c.comment]
        + [f"{c}{lines.max_row}" for c in "BCDEFGHIJKLM"],
        "Findings": [f"{c}{r}" for r in range(1, findings.max_row + 1) for c in "BDEFJ"],
    }
    return {
        "sheets": wb.sheetnames,
        "freeze_panes": {ws.title: ws.freeze_panes for ws in wb},
        "column_widths": {"Line Review B": lines.column_dimensions["B"].width},
        "cells": {
            sheet: {coord: _cell(wb[sheet], coord) for coord in coords}
            for sheet, coords in keys.items()
        },
        "summary_rows": summary.max_row,
    }


@pytest.mark.parametrize(("sub", "app"), [("mech", 4), ("elec", 3), ("dryw", 1)])
def test_workbook_matches_golden(tmp_path: Path, sub: str, app: int) -> None:
    review, result = review_application(SAMPLE, sub, app)
    path = tmp_path / "review.xlsx"
    write_workbook(review, result, path)
    actual = snapshot(path)

    golden = GOLDEN / f"{sub}_app_{app}_workbook.json"
    if os.environ.get("UPDATE_GOLDEN"):
        golden.parent.mkdir(exist_ok=True)
        golden.write_text(json.dumps(actual, indent=2) + "\n", encoding="utf-8", newline="\n")
    assert actual == json.loads(golden.read_text(encoding="utf-8"))
